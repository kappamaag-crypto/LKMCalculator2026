"""Профессиональный экспорт расчёта в Excel (openpyxl)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.models import SystemCalculationResult, ComparisonResult, ObjectData, RecommendationResult
from app.config import AppSettings

HEADER_FONT = Font(name="Arial", bold=True, size=12, color="FFFFFF")
TITLE_FONT = Font(name="Arial", bold=True, size=16)
NORMAL_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", bold=True, size=10)
THIN = Border(
    left=Side(style="thin", color="D0D4DC"), right=Side(style="thin", color="D0D4DC"),
    top=Side(style="thin", color="D0D4DC"), bottom=Side(style="thin", color="D0D4DC"),
)
HEADER_FILL = PatternFill("solid", fgColor="1A56DB")
TOTAL_FILL = PatternFill("solid", fgColor="DBEAFE")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _auto_width(ws: Worksheet, min_w: int = 10, max_w: int = 40) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max((len(str(c.value or "")) for c in col), default=min_w)
        ws.column_dimensions[letter].width = min(max(length + 2, min_w), max_w)


def _header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN


def _cell(ws, row, col, value, bold=False, fill=None, align=CENTER):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = BOLD_FONT if bold else NORMAL_FONT
    cell.alignment = align
    cell.border = THIN
    if fill:
        cell.fill = fill
    return cell


class ExcelExporter:
    def __init__(self, settings: Optional[AppSettings] = None):
        self.settings = settings or AppSettings()

    def export_calculation(
        self, result: SystemCalculationResult, path: str | Path,
        recommendation: Optional[RecommendationResult] = None,
    ) -> Path:
        path = Path(path)
        wb = Workbook()
        self._sheet_input(wb.active, result)
        self._sheet_layers(wb.create_sheet("\u0421\u043b\u043e\u0438"), result)
        self._sheet_materials(wb.create_sheet("\u041c\u0430\u0442\u0435\u0440\u0438\u0430\u043b\u044b"), result)
        self._sheet_summary(wb.create_sheet("\u0418\u0442\u043e\u0433\u0438"), result)
        if recommendation and recommendation.items:
            self._sheet_recommendations(wb.create_sheet("\u0420\u0435\u043a\u043e\u043c\u0435\u043d\u0434\u0430\u0446\u0438\u0438"), recommendation)
        wb.save(path)
        return path

    def export_comparison(self, comparison: ComparisonResult, path: str | Path) -> Path:
        path = Path(path)
        wb = Workbook()
        self._sheet_comparison(wb.active, comparison)
        wb.save(path)
        return path

    def _sheet_input(self, ws: Worksheet, result: SystemCalculationResult) -> None:
        ws.title = "\u0418\u0441\u0445\u043e\u0434\u043d\u044b\u0435 \u0434\u0430\u043d\u043d\u044b\u0435"
        obj = result.object_data
        ws["A1"] = "\u041a\u0430\u043b\u044c\u043a\u0443\u043b\u044f\u0442\u043e\u0440 \u041b\u041a\u041c / \u0410\u041a\u0417 — О\u0442\u0447\u0451\u0442 р\u0430\u0441\u0447\u0451\u0442\u0430"
        ws["A1"].font = TITLE_FONT
        rows = [
            ("\u041e\u0431\u044a\u0435\u043a\u0442", obj.object_name),
            ("\u0417\u0430\u043a\u0430\u0437\u0447\u0438\u043a", obj.customer),
            ("\u041f\u0440\u043e\u0435\u043a\u0442", obj.project),
            ("\u2116 р\u0430\u0441\u0447\u0451\u0442\u0430", obj.calculation_number),
            ("\u041f\u043b\u043e\u0449\u0430\u0434\u044c, \u043c\u00b2", obj.area_m2),
            ("\u041a\u0430\u0442\u0435\u0433\u043e\u0440\u0438\u044f", str(obj.corrosion_category or "—")),
            ("\u0414\u043e\u043b\u0433\u043e\u0432\u0435\u0447\u043d\u043e\u0441\u0442\u044c", str(obj.durability or "—")),
            ("\u0421\u0438\u0441\u0442\u0435\u043c\u0430", result.system.system_name),
            ("\u0414\u0430\u0442\u0430", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ]
        for i, (k, v) in enumerate(rows, 3):
            _cell(ws, i, 1, k, bold=True, align=LEFT)
            _cell(ws, i, 2, v, align=LEFT)
        _auto_width(ws)

    def _sheet_layers(self, ws: Worksheet, result: SystemCalculationResult) -> None:
        headers = ["\u2116", "\u041c\u0430\u0442\u0435\u0440\u0438\u0430\u043b", "\u0421\u0432\u044f\u0437\u0443\u044e\u0449\u0435\u0435", "DFT, \u043c\u043a\u043c", "WFT, \u043c\u043a\u043c",
                   "\u0420\u0430\u0441\u0445\u043e\u0434, \u043a\u0433/\u043c\u00b2", "\u0420\u0430\u0441\u0445\u043e\u0434, \u043b/\u043c\u00b2", "\u0421\u0442\u043e\u0438\u043c\u043e\u0441\u0442\u044c, \u0440\u0443\u0431/\u043c\u00b2",
                   "\u0418\u0442\u043e\u0433\u043e \u043a\u0433", "\u0418\u0442\u043e\u0433\u043e \u0440\u0443\u0431"]
        _header_row(ws, 1, headers)
        for i, lr in enumerate(result.layers, 1):
            binder = lr.material.binder_type.value if hasattr(lr.material.binder_type, "value") else str(lr.material.binder_type)
            vals = [i, lr.material.material_name, binder, lr.target_dft, lr.wft,
                    lr.practical_consumption_kg, lr.practical_consumption_l, lr.cost_per_m2,
                    lr.total_consumption_kg, lr.total_cost]
            for c, v in enumerate(vals, 1):
                _cell(ws, i + 1, c, v if not isinstance(v, float) else round(v, 3))
        total_row = len(result.layers) + 2
        _cell(ws, total_row, 1, "\u0418\u0422\u041e\u0413\u041e", bold=True, fill=TOTAL_FILL)
        _cell(ws, total_row, 4, result.total_dft, bold=True, fill=TOTAL_FILL)
        _cell(ws, total_row, 6, result.total_practical_consumption_kg, bold=True, fill=TOTAL_FILL)
        _cell(ws, total_row, 8, result.total_cost_per_m2, bold=True, fill=TOTAL_FILL)
        _cell(ws, total_row, 10, result.total_cost, bold=True, fill=TOTAL_FILL)
        _auto_width(ws)

    def _sheet_materials(self, ws: Worksheet, result: SystemCalculationResult) -> None:
        headers = ["\u041c\u0430\u0442\u0435\u0440\u0438\u0430\u043b", "\u041f\u043b\u043e\u0442\u043d\u043e\u0441\u0442\u044c", "\u0421\u041e, %", "\u0426\u0435\u043d\u0430, \u0440\u0443\u0431/\u043a\u0433", "\u0424\u0430\u0441\u043e\u0432\u043a\u0430, \u043a\u0433",
                   "\u041d\u0443\u0436\u043d\u043e, \u043a\u0433", "\u0423\u043f\u0430\u043a\u043e\u0432\u043e\u043a", "\u0417\u0430\u043a\u0443\u043f\u043a\u0430, \u043a\u0433"]
        _header_row(ws, 1, headers)
        for i, lr in enumerate(result.layers, 1):
            m = lr.material
            vals = [m.material_name, m.density, m.solids_percent, m.price_per_kg or 0,
                    m.packaging_kg or 0, lr.total_consumption_kg, lr.packages_count, lr.purchase_kg]
            for c, v in enumerate(vals, 1):
                _cell(ws, i + 1, c, v if not isinstance(v, float) else round(v, 3))
        _auto_width(ws)

    def _sheet_summary(self, ws: Worksheet, result: SystemCalculationResult) -> None:
        ws["A1"] = "\u0418\u0442\u043e\u0433\u0438 р\u0430\u0441\u0447\u0451\u0442\u0430"
        ws["A1"].font = TITLE_FONT
        rows = [
            ("\u0421\u043b\u043e\u0451\u0432", len(result.layers)),
            ("\u041e\u0431\u0449\u0430\u044f DFT, \u043c\u043a\u043c", result.total_dft),
            ("\u0420\u0430\u0441\u0445\u043e\u0434, \u043a\u0433/\u043c\u00b2", result.total_practical_consumption_kg),
            ("\u0420\u0430\u0441\u0445\u043e\u0434, \u043b/\u043c\u00b2", result.total_practical_consumption_l),
            ("\u0421\u0442\u043e\u0438\u043c\u043e\u0441\u0442\u044c, \u0440\u0443\u0431/\u043c\u00b2", result.total_cost_per_m2),
            ("\u0421\u0442\u043e\u0438\u043c\u043e\u0441\u0442\u044c \u043e\u0431\u044a\u0435\u043a\u0442\u0430, \u0440\u0443\u0431", result.total_cost),
            ("\u0421\u0442\u043e\u0438\u043c\u043e\u0441\u0442\u044c \u0437\u0430\u043a\u0443\u043f\u043a\u0438, \u0440\u0443\u0431", result.total_purchase_cost),
        ]
        for i, (k, v) in enumerate(rows, 3):
            _cell(ws, i, 1, k, bold=True, align=LEFT)
            _cell(ws, i, 2, round(v, 3) if isinstance(v, float) else v)
        _auto_width(ws)

    def _sheet_comparison(self, ws: Worksheet, comparison: ComparisonResult) -> None:
        ws.title = "\u0421\u0440\u0430\u0432\u043d\u0435\u043d\u0438\u0435"
        systems = comparison.systems
        headers = ["\u041f\u043e\u043a\u0430\u0437\u0430\u0442\u0435\u043b\u044c"] + [(s.system.system_name or f"S{i+1}")[:25] for i, s in enumerate(systems)]
        _header_row(ws, 1, headers)
        rows_data = [
            ("\u0421\u043b\u043e\u0451\u0432", [len(s.layers) for s in systems]),
            ("DFT, \u043c\u043a\u043c", [s.total_dft for s in systems]),
            ("\u0420\u0430\u0441\u0445\u043e\u0434, \u043a\u0433/\u043c\u00b2", [s.total_practical_consumption_kg for s in systems]),
            ("\u0421\u0442\u043e\u0438\u043c\u043e\u0441\u0442\u044c, \u0440\u0443\u0431/\u043c\u00b2", [s.total_cost_per_m2 for s in systems]),
            ("\u041e\u0431\u044a\u0435\u043a\u0442, \u0440\u0443\u0431", [s.total_cost for s in systems]),
        ]
        for r, (label, values) in enumerate(rows_data, 2):
            _cell(ws, r, 1, label, bold=True, align=LEFT)
            for c, v in enumerate(values, 2):
                _cell(ws, r, c, round(v, 2) if isinstance(v, float) else v)
        _auto_width(ws)

    def _sheet_recommendations(self, ws: Worksheet, rec: RecommendationResult) -> None:
        headers = ["\u041c\u0435\u0441\u0442\u043e", "\u0421\u0438\u0441\u0442\u0435\u043c\u0430", "Score", "\u041f\u0440\u0438\u0447\u0438\u043d\u044b"]
        _header_row(ws, 1, headers)
        for i, item in enumerate(rec.items, 1):
            reasons = "; ".join(item.reasons[:3])
            for c, v in enumerate([item.rank, item.system.system_name, item.score, reasons], 1):
                _cell(ws, i + 1, c, v, align=LEFT if c in (2, 4) else CENTER)
        _auto_width(ws)
