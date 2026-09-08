"""Тесты Excel-экспорта."""

from __future__ import annotations

from pathlib import Path
import tempfile

import pytest
from openpyxl import load_workbook

from app.domain.models import Material, ObjectData
from app.domain.enums import MaterialType, BinderType, CorrosionCategory, DurabilityLevel
from app.domain.calculator import LayerInput
from app.services.calculation_service import CalculationService
from app.infrastructure.export.excel_exporter import ExcelExporter


@pytest.fixture
def sample_result():
    primer = Material(
        material_name="Грунт-Эмаль Blank Universal",
        material_type=MaterialType.PRIMER_ENAMEL,
        binder_type=BinderType.EPOXY,
        density=1.4, solids_percent=73.0, price_per_kg=552.0, packaging_kg=20.0,
        manufacturer="Blank",
    )
    finish = Material(
        material_name="Эмаль Blank Finish",
        material_type=MaterialType.FINISH,
        binder_type=BinderType.POLYURETHANE,
        density=1.3, solids_percent=58.0, price_per_kg=892.0, packaging_kg=20.0,
        manufacturer="Blank",
    )
    obj = ObjectData(
        object_name="Тест", customer="Заказчик",
        area_m2=100.0, calculation_number="T-001",
        corrosion_category=CorrosionCategory.C4,
        durability=DurabilityLevel.HIGH,
    )
    layers = [
        LayerInput(material=primer, target_dft=200, losses_percent=0),
        LayerInput(material=finish, target_dft=100, losses_percent=0),
    ]
    result, validation = CalculationService().calculate_system(obj, layers)
    assert not validation.has_errors
    return result


def test_export_creates_file(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"
        ExcelExporter().export_calculation(sample_result, path)
        assert path.exists()
        assert path.stat().st_size > 1000


def test_export_has_required_sheets(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"
        ExcelExporter().export_calculation(sample_result, path)
        wb = load_workbook(path)
        names = set(wb.sheetnames)
        assert "Исходные данные" in names
        assert "Слои" in names
        assert "Материалы" in names
        assert "Итоги" in names


def test_layers_sheet_has_data(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"
        ExcelExporter().export_calculation(sample_result, path)
        wb = load_workbook(path)
        ws = wb["Слои"]
        # Должны быть заголовки и минимум 2 строки данных + итого
        assert ws.max_row >= 4
        # Проверяем наличие названия материала
        found = False
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=3):
            for cell in row:
                if cell.value and "Blank" in str(cell.value):
                    found = True
        assert found


def test_export_comparison(sample_result):
    from app.domain.comparison import ComparisonEngine
    from app.domain.calculator import LayerInput

    primer = sample_result.layers[0].material
    finish = sample_result.layers[1].material
    obj = sample_result.object_data

    sys1 = [
        LayerInput(material=primer, target_dft=200),
        LayerInput(material=finish, target_dft=100),
    ]
    sys2 = [
        LayerInput(material=primer, target_dft=150),
        LayerInput(material=finish, target_dft=80),
    ]
    comparison = ComparisonEngine().compare(obj, [("A", sys1), ("B", sys2)])

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "cmp.xlsx"
        ExcelExporter().export_comparison(comparison, path)
        assert path.exists()
        wb = load_workbook(path)
        assert "Сравнение" in wb.sheetnames
