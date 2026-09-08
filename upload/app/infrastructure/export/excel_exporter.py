"""Профессиональный экспорт расчёта / сравнения в Excel."""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from app.domain.models import SystemCalculationResult, ComparisonResult, ObjectData, RecommendationResult
from app.config import AppSettings

HEADER_FONT=Font(name="Arial",bold=True,size=12,color="FFFFFF")
TITLE_FONT=Font(name="Arial",bold=True,size=16,color="1A1A2E")
SUBTITLE_FONT=Font(name="Arial",bold=True,size=11,color="374151")
NORMAL_FONT=Font(name="Arial",size=10)
BOLD_FONT=Font(name="Arial",bold=True,size=10)
THIN=Border(left=Side(style="thin",color="D0D4DC"),right=Side(style="thin",color="D0D4DC"),top=Side(style="thin",color="D0D4DC"),bottom=Side(style="thin",color="D0D4DC"))
HEADER_FILL=PatternFill("solid",fgColor="1A56DB")
ALT_FILL=PatternFill("solid",fgColor="F3F4F6")
TOTAL_FILL=PatternFill("solid",fgColor="DBEAFE")
CENTER=Alignment(horizontal="center",vertical="center",wrap_text=True)
LEFT=Alignment(horizontal="left",vertical="center",wrap_text=True)


def _auto_width(ws:Worksheet,min_width:int=10,max_width:int=40)->None:
    for col in ws.columns:
        letter=get_column_letter(col[0].column)
        length=max((len(str(c.value or "")) for c in col),default=min_width)
        ws.column_dimensions[letter].width=min(max(length+2,min_width),max_width)


def _write_header_row(ws:Worksheet,row:int,headers:list[str])->None:
    for col,header in enumerate(headers,1):
        cell=ws.cell(row=row,column=col,value=header)
        cell.font=HEADER_FONT; cell.fill=HEADER_FILL; cell.alignment=CENTER; cell.border=THIN


def _cell(ws:Worksheet,row:int,col:int,value,bold:bool=False,fill=None,align=CENTER):
    cell=ws.cell(row=row,column=col,value=value)
    cell.font=BOLD_FONT if bold else NORMAL_FONT; cell.alignment=align; cell.border=THIN
    if fill: cell.fill=fill
    return cell


def _area(obj:ObjectData)->float:
    """Площадь расчёта задаётся непосредственно пользователем, в м²."""
    return max(float(obj.area_m2 or 0.0), 0.0)


def _cost_add(*values: Optional[float]) -> Optional[float]:
    if any(value is None for value in values): return None
    return sum(values)


def _display(value, digits:int=2):
    return "—" if value is None else round(value, digits)


class ExcelExporter:
    """Экспорт инженерного расчёта без закупочной и складской логики."""
    def __init__(self,settings:Optional[AppSettings]=None): self.settings=settings or AppSettings()

    def export_calculation(self,result:SystemCalculationResult,path:str|Path,recommendation:Optional[RecommendationResult]=None)->Path:
        path=Path(path); wb=Workbook(); self._sheet_input(wb.active,result); self._sheet_layers(wb.create_sheet("Слои"),result); self._sheet_materials(wb.create_sheet("Материалы"),result); self._sheet_summary(wb.create_sheet("Итоги"),result)
        if recommendation and recommendation.items: self._sheet_recommendations(wb.create_sheet("Рекомендации"),recommendation)
        wb.save(path); return path

    def export_comparison(self,comparison:ComparisonResult,path:str|Path)->Path:
        path=Path(path); wb=Workbook(); ws=wb.active; ws.title="Сравнение"; self._sheet_comparison(ws,comparison)
        for i,result in enumerate(comparison.systems): self._sheet_layers(wb.create_sheet((result.system.system_name or f"Система {i+1}")[:28]),result)
        wb.save(path); return path

    def _title_block(self,ws:Worksheet,title:str,obj:ObjectData,start_row:int=1)->int:
        ws.cell(start_row,1,title).font=TITLE_FONT; ws.merge_cells(start_row=start_row,start_column=1,end_row=start_row,end_column=6); row=start_row+1
        ws.cell(row,1,self.settings.organization_name).font=SUBTITLE_FONT; row+=1; ws.cell(row,1,f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}").font=NORMAL_FONT; row+=1
        for label,value in (("№ расчёта",obj.calculation_number),("Объект",obj.object_name),("Заказчик",obj.customer)):
            if value: ws.cell(row,1,f"{label}: {value}").font=NORMAL_FONT; row+=1
        ws.cell(row,1,f"Площадь: {_area(obj):.2f} м²").font=NORMAL_FONT; return row+2

    def _sheet_input(self,ws:Worksheet,result:SystemCalculationResult)->None:
        ws.title="Исходные данные"; row=self._title_block(ws,"Расчёт расхода ЛКМ / подбор системы АКЗ",result.object_data); ws.cell(row,1,"Параметры расчёта").font=SUBTITLE_FONT; row+=1
        fields=[("Категория коррозии",result.object_data.corrosion_category.value if result.object_data.corrosion_category else "—"),("Долговечность",result.object_data.durability.value if result.object_data.durability else "—"),("Поверхность",result.object_data.surface_type.value if result.object_data.surface_type else "—"),("Среда",result.object_data.environment.value if result.object_data.environment else "—"),("Система",result.system.system_name or "Пользовательская")]
        for label,value in fields: _cell(ws,row,1,label,True,align=LEFT); _cell(ws,row,2,value,align=LEFT); row+=1
        _auto_width(ws)

    def _sheet_layers(self,ws:Worksheet,result:SystemCalculationResult)->None:
        ws.title="Слои" if ws.title=="Sheet" else ws.title; ws.cell(1,1,f"Слои системы: {result.system.system_name}").font=TITLE_FONT; row=3
        headers=["№","Материал","Связующее","DFT, мкм","WFT, мкм","Потери, %","Разб., %","Укрыв. теор., м²/л","Укрыв. практ., м²/л","Расход теор., кг/м²","Расход практ., кг/м²","Расход практ., л/м²","Разбавитель, кг/м²","Разбавитель, л/м²","Стоимость материала, руб/м²","Разбавитель, руб/м²","Итого, руб/м²","Расход ЛКМ на объект, кг","Разбавитель на объект, кг","Стоимость на объект, руб"]
        _write_header_row(ws,row,headers); row+=1
        for i,lr in enumerate(result.layers):
            fill=ALT_FILL if i%2 else None; binder=lr.material.binder_type.value if hasattr(lr.material.binder_type,"value") else str(lr.material.binder_type); total_m2=_cost_add(lr.cost_per_m2,lr.thinner_cost_per_m2); area=_area(result.object_data)
            thinner_total_kg=lr.thinner_consumption_kg*area
            values=[i+1,lr.material.material_name,binder,lr.target_dft,lr.wft,lr.losses_percent,lr.thinner_percent,lr.theoretical_coverage,lr.practical_coverage,lr.theoretical_consumption_kg,lr.practical_consumption_kg,lr.practical_consumption_l,lr.thinner_consumption_kg,lr.thinner_consumption_l,_display(lr.cost_per_m2),_display(lr.thinner_cost_per_m2),_display(total_m2),lr.total_consumption_kg,thinner_total_kg,_display(lr.total_cost)]
            for col,value in enumerate(values,1): _cell(ws,row,col,value,fill=fill)
            row+=1
        area=_area(result.object_data); thinner_cost_m2=result.total_thinner_cost/area if area and result.total_thinner_cost is not None else None; material_cost=_cost_add(result.total_cost_per_m2, -thinner_cost_m2) if result.total_cost_per_m2 is not None and thinner_cost_m2 is not None else None; thinner_kg_m2=sum(lr.thinner_consumption_kg for lr in result.layers); thinner_l_m2=sum(lr.thinner_consumption_l for lr in result.layers)
        _cell(ws,row,2,"ИТОГО",True,TOTAL_FILL,LEFT); _cell(ws,row,4,result.total_dft,True,TOTAL_FILL); _cell(ws,row,11,result.total_practical_consumption_kg,True,TOTAL_FILL); _cell(ws,row,12,result.total_practical_consumption_l,True,TOTAL_FILL); _cell(ws,row,13,thinner_kg_m2,True,TOTAL_FILL); _cell(ws,row,14,thinner_l_m2,True,TOTAL_FILL); _cell(ws,row,15,_display(material_cost),True,TOTAL_FILL); _cell(ws,row,16,_display(thinner_cost_m2),True,TOTAL_FILL); _cell(ws,row,17,_display(result.total_cost_per_m2),True,TOTAL_FILL); _cell(ws,row,18,result.total_practical_consumption_kg*area,True,TOTAL_FILL); _cell(ws,row,19,thinner_kg_m2*area,True,TOTAL_FILL); _cell(ws,row,20,_display(result.total_cost),True,TOTAL_FILL); ws.freeze_panes="A4"; _auto_width(ws,8,28)

    def _sheet_materials(self,ws:Worksheet,result:SystemCalculationResult)->None:
        ws.cell(1,1,"Спецификация материалов").font=TITLE_FONT; row=3; _write_header_row(ws,row,["Материал","Производитель","Плотность, кг/л","Сухой остаток, %","Цена, руб/кг","Цена, руб/л","Расход ЛКМ, кг","Расход ЛКМ, л","Расход разбавителя, кг","Расход разбавителя, л","Стоимость материала, руб","Разбавитель, руб"]); row+=1; area=_area(result.object_data)
        for i,lr in enumerate(result.layers):
            fill=ALT_FILL if i%2 else None; material_total=lr.cost_per_m2*area if lr.cost_per_m2 is not None else None; thinner_total=lr.thinner_cost_per_m2*area if lr.thinner_cost_per_m2 is not None else None
            values=[lr.material.material_name,lr.material.manufacturer or "—",_display(lr.material.density),_display(lr.material.solids_by_volume_percent if lr.material.solids_by_volume_percent is not None else lr.material.solids_percent),_display(lr.material.price_per_kg),_display(lr.material.price_per_liter),lr.total_consumption_kg,lr.total_consumption_l,lr.thinner_consumption_kg*area,lr.thinner_consumption_l*area,_display(material_total),_display(thinner_total)]
            for col,value in enumerate(values,1): _cell(ws,row,col,value,fill=fill)
            row+=1
        _auto_width(ws)

    def _sheet_summary(self,ws:Worksheet,result:SystemCalculationResult)->None:
        row=self._title_block(ws,"Итоговый расчёт",result.object_data); ws.cell(row,1,"Сводка").font=SUBTITLE_FONT; row+=1; area=_area(result.object_data); thinner_cost_m2=result.total_thinner_cost/area if area and result.total_thinner_cost is not None else None; material_cost_m2=_cost_add(result.total_cost_per_m2,-thinner_cost_m2) if result.total_cost_per_m2 is not None and thinner_cost_m2 is not None else None; thinner_l_m2=sum(lr.thinner_consumption_l for lr in result.layers); thinner_kg_m2=sum(lr.thinner_consumption_kg for lr in result.layers)
        items=[("Система",result.system.system_name or "Пользовательская"),("Количество слоёв",len(result.layers)),("Общая толщина DFT, мкм",result.total_dft),("Расход ЛКМ теоретический, кг/м²",result.total_theoretical_consumption_kg),("Расход ЛКМ практический, кг/м²",result.total_practical_consumption_kg),("Расход ЛКМ практический, л/м²",result.total_practical_consumption_l),("Расход разбавителя, кг/м²",thinner_kg_m2),("Расход разбавителя, л/м²",thinner_l_m2),("Стоимость ЛКМ, руб/м²",_display(material_cost_m2)),("Стоимость разбавителя, руб/м²",_display(thinner_cost_m2)),("Итого, руб/м²",_display(result.total_cost_per_m2)),("Стоимость объекта, руб",_display(result.total_cost)),("В т.ч. разбавитель, руб",_display(result.total_thinner_cost))]
        for label,value in items: _cell(ws,row,1,label,True,ALT_FILL,LEFT); _cell(ws,row,2,value,align=LEFT); row+=1
        _auto_width(ws)

    def _sheet_comparison(self,ws:Worksheet,comparison:ComparisonResult)->None:
        row=self._title_block(ws,"Сравнение систем покрытия",comparison.object_data); systems=comparison.systems; _write_header_row(ws,row,["Показатель"]+[(s.system.system_name or f"Система {i+1}")[:25] for i,s in enumerate(systems)]); row+=1
        rows=[("Количество слоёв",[len(s.layers) for s in systems]),("Общая толщина, мкм",[s.total_dft for s in systems]),("Практический расход ЛКМ, кг/м²",[s.total_practical_consumption_kg for s in systems]),("Практический расход ЛКМ, л/м²",[s.total_practical_consumption_l for s in systems]),("Стоимость, руб/м²",[_display(s.total_cost_per_m2) for s in systems]),("Стоимость объекта, руб",[_display(s.total_cost,0) for s in systems])]
        for label,values in rows:
            _cell(ws,row,1,label,True,align=LEFT)
            for col,value in enumerate(values,2): _cell(ws,row,col,value)
            row+=1
        _auto_width(ws)

    def _sheet_recommendations(self,ws:Worksheet,recommendation:RecommendationResult)->None:
        ws.cell(1,1,"Рекомендации по системам").font=TITLE_FONT; row=3; _write_header_row(ws,row,["Место","Система","Статус","Баллы","Причины / ограничения"]); row+=1
        for item in recommendation.items:
            reasons="; ".join(item.reasons+item.warnings+item.limitations)
            for col,value in enumerate([item.rank,item.system.system_name,item.status,item.score,reasons],1): _cell(ws,row,col,value,align=LEFT if col in (2,5) else CENTER)
            row+=1
        _auto_width(ws)
