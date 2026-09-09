"""Экспорт расчёта в пользовательский Excel-шаблон.

Шаблон сохраняется, но инженерные данные расчёта всегда записываются
в отдельный лист «Расчёт» и, если возможно, также в найденные поля шаблона.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import AppSettings
from app.domain.formulas import FORMULA_VERSION
from app.domain.models import SystemCalculationResult, RecommendationResult
from app.infrastructure.export.excel_exporter import ExcelExporter


class CustomerExcelExporter(ExcelExporter):
    LABELS = {
        "object": ("объект", "название объекта"),
        "customer": ("заказчик",),
        "area": ("площадь", "площадь, м²", "площадь м2"),
        "system": ("система", "система покрытия"),
        "corrosion": ("категория коррозии", "категория"),
        "durability": ("долговечность",),
        "total_dft": ("общая толщина", "суммарная толщина", "толщина dft"),
        "cost_m2": ("стоимость, руб/м²", "стоимость руб/м2", "стоимость м²"),
        "total_cost": ("стоимость объекта", "итого стоимость"),
        "calculation_date": ("дата расчёта",),
        "formula_version": ("версия формул", "версия расчёта"),
    }

    def __init__(self, settings: Optional[AppSettings] = None):
        super().__init__(settings)

    @staticmethod
    def _text(value) -> str:
        return "" if value is None else " ".join(str(value).strip().lower().replace("ё", "е").split())

    def _find_label_cell(self, wb, aliases):
        aliases = tuple(self._text(a) for a in aliases)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    text = self._text(cell.value)
                    if text and any(alias in text for alias in aliases):
                        return ws, cell
        return None, None

    def _put_next_to_label(self, wb, aliases, value) -> bool:
        ws, cell = self._find_label_cell(wb, aliases)
        if cell is None:
            return False
        for offset in (1, 2, 3):
            target = ws.cell(cell.row, cell.column + offset)
            if target.value in (None, ""):
                target.value = value
                return True
        return False

    @staticmethod
    def _cost_add(*values):
        if any(value is None for value in values):
            return None
        return sum(values)

    @staticmethod
    def _area(result) -> float:
        return max(float(result.object_data.area_m2 or 0.0), 0.0)

    def _write_engineering_sheet(self, wb, result: SystemCalculationResult) -> None:
        """Гарантированно выгружает именно выбранные пользователем слои."""
        name = "Расчёт"
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A5"

        title_font = Font(name="Arial", bold=True, size=15)
        header_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        normal_font = Font(name="Arial", size=10)
        fill = PatternFill("solid", fgColor="1A56DB")
        border = Border(left=Side(style="thin", color="D0D4DC"), right=Side(style="thin", color="D0D4DC"), top=Side(style="thin", color="D0D4DC"), bottom=Side(style="thin", color="D0D4DC"))

        obj = result.object_data
        ws.cell(1, 1, f"Расчёт системы покрытия: {result.system.system_name or 'Пользовательская система'}").font = title_font
        ws.cell(2, 1, f"Объект: {obj.object_name or '—'}")
        ws.cell(2, 4, f"Заказчик: {obj.customer or '—'}")
        ws.cell(3, 1, f"Площадь: {obj.area_m2:.2f} м²" if obj.area_m2 else "Площадь: —")
        ws.cell(3, 4, f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        ws.cell(4, 1, f"Формула: {FORMULA_VERSION}")

        headers = [
            "№", "Материал", "Связующее", "DFT, мкм", "WFT, мкм",
            "Потери, %", "Разбавитель, %", "Расход, кг/м²", "Расход, л/м²",
            "Разбавитель, кг/м²", "Разбавитель, л/м²", "Стоимость, руб/м²",
            "Стоимость на объект, руб"
        ]
        header_row = 6
        for col, value in enumerate(headers, 1):
            c = ws.cell(header_row, col, value)
            c.font = header_font
            c.fill = fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border

        area = self._area(result)
        for i, layer in enumerate(result.layers, 1):
            binder = getattr(layer.material.binder_type, "value", layer.material.binder_type)
            total_cost_m2 = self._cost_add(layer.cost_per_m2, layer.thinner_cost_per_m2)
            values = [
                i,
                layer.material.material_name,
                binder,
                layer.target_dft,
                layer.wft,
                layer.losses_percent,
                layer.thinner_percent,
                layer.practical_consumption_kg,
                layer.practical_consumption_l,
                layer.thinner_consumption_kg,
                layer.thinner_consumption_l,
                total_cost_m2,
                layer.total_cost,
            ]
            row = header_row + i
            for col, value in enumerate(values, 1):
                c = ws.cell(row, col, value)
                c.font = normal_font
                c.border = border
                c.alignment = Alignment(horizontal="left" if col == 2 else "center", vertical="center", wrap_text=True)

        total_row = header_row + len(result.layers) + 1
        ws.cell(total_row, 2, "ИТОГО").font = Font(name="Arial", bold=True, size=10)
        ws.cell(total_row, 4, result.total_dft)
        ws.cell(total_row, 8, result.total_practical_consumption_kg)
        ws.cell(total_row, 9, result.total_practical_consumption_l)
        ws.cell(total_row, 12, result.total_cost_per_m2)
        ws.cell(total_row, 13, result.total_cost)
        for col in (2, 4, 8, 9, 12, 13):
            ws.cell(total_row, col).border = border
            ws.cell(total_row, col).font = Font(name="Arial", bold=True, size=10)

        for col in range(1, len(headers) + 1):
            max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, total_row + 1))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 32)

    def export_calculation(self, result: SystemCalculationResult, path: str | Path, recommendation: Optional[RecommendationResult] = None) -> Path:
        path = Path(path)
        template = Path(self.settings.excel_template_path)

        # Если шаблон недоступен — используем штатный профессиональный экспорт.
        if not template.exists() or template.resolve() == path.resolve():
            return super().export_calculation(result, path, recommendation)

        wb = load_workbook(template)
        obj = result.object_data
        values = {
            "object": obj.object_name or None,
            "customer": obj.customer or None,
            "area": obj.area_m2 if obj.area_m2 > 0 else None,
            "system": result.system.system_name or None,
            "corrosion": obj.corrosion_category.value if obj.corrosion_category else None,
            "durability": obj.durability.value if obj.durability else None,
            "total_dft": result.total_dft,
            "cost_m2": result.total_cost_per_m2,
            "total_cost": result.total_cost,
            "calculation_date": datetime.now().strftime("%d.%m.%Y %H:%M"),
            "formula_version": FORMULA_VERSION,
        }
        for key, value in values.items():
            if value is not None:
                self._put_next_to_label(wb, self.LABELS[key], value)

        # Ключевое исправление: выбранные слои больше не зависят от структуры шаблона.
        # Они всегда попадают в отдельный лист «Расчёт».
        self._write_engineering_sheet(wb, result)
        wb.save(path)
        return path
