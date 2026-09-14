"""Инженерный Excel-экспорт систем покрытия без складских и закупочных данных."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.models import ComparisonResult, SystemCalculationResult


class EngineeringExcelExporter:
    """Экспортирует технический результат в одну компактную книгу Excel."""

    _HEADER_FILL = PatternFill("solid", fgColor="1A56DB")
    _TOTAL_FILL = PatternFill("solid", fgColor="DBEAFE")
    _BORDER = Border(
        left=Side(style="thin", color="D0D4DC"),
        right=Side(style="thin", color="D0D4DC"),
        top=Side(style="thin", color="D0D4DC"),
        bottom=Side(style="thin", color="D0D4DC"),
    )
    _HEADER_FONT = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    _TITLE_FONT = Font(name="Arial", bold=True, size=15, color="1A1A2E")
    _NORMAL_FONT = Font(name="Arial", size=10)
    _BOLD_FONT = Font(name="Arial", bold=True, size=10)

    @staticmethod
    def _area(result: SystemCalculationResult) -> float:
        return max(float(result.object_data.area_m2 or 0.0), 0.0)

    @staticmethod
    def _money(value: Optional[float]):
        return "—" if value is None else round(value, 2)

    def _cell(self, ws, row: int, col: int, value, *, bold=False, fill=None, left=False):
        cell = ws.cell(row, col, value)
        cell.font = self._BOLD_FONT if bold else self._NORMAL_FONT
        cell.alignment = Alignment(horizontal="left" if left else "center", vertical="center", wrap_text=True)
        cell.border = self._BORDER
        if fill:
            cell.fill = fill
        return cell

    def _header(self, ws, row: int, headers: list[str]) -> None:
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row, col, header)
            cell.font = self._HEADER_FONT
            cell.fill = self._HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self._BORDER
        ws.row_dimensions[row].height = 38

    def _write_system_table(self, ws, result: SystemCalculationResult, start_row: int) -> int:
        headers = [
            "Слой", "Материал", "Связующее", "DFT, мкм", "WFT, мкм",
            "Расход теор., кг/м²", "Расход практ., кг/м²", "Расход практ., л/м²",
            "Разбавитель", "Разбавитель, кг/м²", "Разбавитель, л/м²",
            "Потери, %", "Стоимость слоя, руб/м²",
        ]
        self._header(ws, start_row, headers)
        row = start_row + 1
        for index, layer in enumerate(result.layers, 1):
            binder = layer.material.binder_type.value if hasattr(layer.material.binder_type, "value") else str(layer.material.binder_type or "—")
            thinner = layer.thinner.display_name() if layer.thinner is not None else "—"
            layer_cost = None
            if layer.cost_per_m2 is not None or layer.thinner_cost_per_m2 is not None:
                layer_cost = (layer.cost_per_m2 or 0.0) + (layer.thinner_cost_per_m2 or 0.0)
            values = [
                index, layer.material.display_name(), binder, layer.target_dft, layer.wft,
                layer.theoretical_consumption_kg, layer.practical_consumption_kg,
                layer.practical_consumption_l, thinner, layer.thinner_consumption_kg,
                layer.thinner_consumption_l, layer.losses_percent, self._money(layer_cost),
            ]
            for col, value in enumerate(values, 1):
                self._cell(ws, row, col, value, left=col in {2, 3, 9})
            row += 1

        thinner_kg = sum(layer.thinner_consumption_kg for layer in result.layers)
        thinner_l = sum(layer.thinner_consumption_l for layer in result.layers)
        self._cell(ws, row, 1, "ИТОГО", bold=True, fill=self._TOTAL_FILL, left=True)
        self._cell(ws, row, 4, result.total_dft, bold=True, fill=self._TOTAL_FILL)
        self._cell(ws, row, 7, result.total_practical_consumption_kg, bold=True, fill=self._TOTAL_FILL)
        self._cell(ws, row, 8, result.total_practical_consumption_l, bold=True, fill=self._TOTAL_FILL)
        self._cell(ws, row, 10, thinner_kg, bold=True, fill=self._TOTAL_FILL)
        self._cell(ws, row, 11, thinner_l, bold=True, fill=self._TOTAL_FILL)
        self._cell(ws, row, 13, self._money(result.total_cost_per_m2), bold=True, fill=self._TOTAL_FILL)
        return row + 3

    def _write_summary_table(self, ws, result: SystemCalculationResult, start_row: int) -> None:
        self._header(ws, start_row, ["Итог системы", "Значение", "Единица"])
        thinner_kg = sum(layer.thinner_consumption_kg for layer in result.layers)
        thinner_l = sum(layer.thinner_consumption_l for layer in result.layers)
        rows = [
            ("Количество слоёв", len(result.layers), "шт."),
            ("Общая толщина", result.total_dft, "мкм"),
            ("Общий практический расход ЛКМ", result.total_practical_consumption_kg, "кг/м²"),
            ("Общий практический расход ЛКМ", result.total_practical_consumption_l, "л/м²"),
            ("Общий расход разбавителя", thinner_kg, "кг/м²"),
            ("Общий расход разбавителя", thinner_l, "л/м²"),
            ("Стоимость системы", self._money(result.total_cost_per_m2), "руб/м²"),
            ("Стоимость объекта", self._money(result.total_cost), "руб"),
            ("Площадь объекта", self._area(result), "м²"),
        ]
        for offset, (label, value, unit) in enumerate(rows, 1):
            self._cell(ws, start_row + offset, 1, label, left=True)
            self._cell(ws, start_row + offset, 2, value)
            self._cell(ws, start_row + offset, 3, unit)

    def export_system(self, result: SystemCalculationResult, path: str | Path) -> Path:
        path = Path(path)
        wb = Workbook()
        ws = wb.active
        ws.title = "Инженерный расчёт"
        ws.sheet_view.showGridLines = False
        title = ws.cell(1, 1, f"Инженерный расчёт системы АКЗ — {result.system.system_name or 'Пользовательская система'}")
        title.font = self._TITLE_FONT
        ws.merge_cells("A1:M1")
        ws.cell(2, 1, f"Объект: {result.object_data.object_name or '—'} | Заказчик: {result.object_data.customer or '—'}")
        ws.merge_cells("A2:M2")
        ws.cell(3, 1, f"Площадь: {self._area(result):.2f} м²")
        ws.merge_cells("A3:M3")
        next_row = self._write_system_table(ws, result, 5)
        self._write_summary_table(ws, result, next_row)
        widths = [8, 34, 16, 12, 12, 20, 22, 20, 30, 22, 22, 12, 24]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A6"
        ws.auto_filter.ref = f"A5:M{next_row - 3}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.save(path)
        return path

    def export_comparison(self, comparison: ComparisonResult, path: str | Path) -> Path:
        """Экспортирует сравнение теми же показателями, которые показаны в UI."""
        if len(comparison.systems) == 1:
            return self.export_system(comparison.systems[0], path)
        if not comparison.systems:
            raise ValueError("Нет систем для экспорта")

        path = Path(path)
        wb = Workbook()
        ws = wb.active
        ws.title = "Сравнение систем"
        ws.sheet_view.showGridLines = False
        system_names = [s.system.system_name or f"Система {i + 1}" for i, s in enumerate(comparison.systems)]
        end_col = max(2, len(system_names) + 1)
        title = ws.cell(1, 1, "Инженерное сравнение систем покрытия")
        title.font = self._TITLE_FONT
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws.cell(2, 1, f"Объект: {comparison.object_data.object_name or '—'}")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        row = 4
        self._header(ws, row, ["Показатель"] + system_names)

        thinner_m2 = [
            (s.total_thinner_cost / s.object_data.area_m2)
            if s.total_thinner_cost is not None and s.object_data.area_m2
            else None
            for s in comparison.systems
        ]
        paint_m2 = [
            (s.total_cost_per_m2 - thinner)
            if s.total_cost_per_m2 is not None and thinner is not None
            else None
            for s, thinner in zip(comparison.systems, thinner_m2)
        ]

        rows = [
            ("Количество слоёв", [len(s.layers) for s in comparison.systems]),
            ("Общая толщина, мкм", [s.total_dft for s in comparison.systems]),
            ("Расход ЛКМ, кг/м²", [s.total_practical_consumption_kg for s in comparison.systems]),
            ("Расход ЛКМ, л/м²", [s.total_practical_consumption_l for s in comparison.systems]),
            ("Стоимость ЛКМ, руб/м²", paint_m2),
            ("Стоимость ЛКМ на объект, руб", [
                (s.total_cost - (s.total_thinner_cost or 0.0)) if s.total_cost is not None else None
                for s in comparison.systems
            ]),
            ("Разбавитель, руб/м²", thinner_m2),
            ("ЛКМ + разбавитель, руб/м²", [s.total_cost_per_m2 for s in comparison.systems]),
            ("Стоимость объекта, руб", [s.total_cost for s in comparison.systems]),
        ]
        technology = [self._technology_values(s) for s in comparison.systems]
        if technology:
            for indicator in technology[0]:
                rows.append((indicator, [item[indicator] for item in technology]))

        max_layers = max(len(s.layers) for s in comparison.systems)
        for index in range(max_layers):
            rows.extend([
                (f"Слой {index + 1}: материал", [
                    s.layers[index].material.display_name() if index < len(s.layers) else "—"
                    for s in comparison.systems
                ]),
                (f"Слой {index + 1}: DFT, мкм", [
                    s.layers[index].target_dft if index < len(s.layers) else None
                    for s in comparison.systems
                ]),
                (f"Слой {index + 1}: WFT, мкм", [
                    s.layers[index].wft if index < len(s.layers) else None
                    for s in comparison.systems
                ]),
                (f"Слой {index + 1}: расход, кг/м²", [
                    s.layers[index].practical_consumption_kg if index < len(s.layers) else None
                    for s in comparison.systems
                ]),
                (f"Слой {index + 1}: расход, л/м²", [
                    s.layers[index].practical_consumption_l if index < len(s.layers) else None
                    for s in comparison.systems
                ]),
                (f"Слой {index + 1}: стоимость, руб/м²", [
                    s.layers[index].cost_per_m2 if index < len(s.layers) else None
                    for s in comparison.systems
                ]),
                (f"Слой {index + 1}: разбавитель, л/м²", [
                    s.layers[index].thinner_consumption_l if index < len(s.layers) else None
                    for s in comparison.systems
                ]),
            ])

        for label, values in rows:
            row += 1
            self._cell(ws, row, 1, label, bold=True, left=True)
            for col, value in enumerate(values, 2):
                formatted = self._money(value) if ("стоимость" in label.lower() or "руб" in label.lower()) else value
                self._cell(ws, row, col, formatted)

        ws.freeze_panes = "B5"
        for col in range(1, ws.max_column + 1):
            max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(col)].width = min(max(18, max_len + 2), 48)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:4"
        ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        wb.save(path)
        return path

    @staticmethod
    def _technology_values(result: SystemCalculationResult) -> dict[str, object]:
        layers = result.layers
        binders = []
        for layer in layers:
            value = layer.material.binder_type.value if hasattr(layer.material.binder_type, "value") else str(layer.material.binder_type or "")
            if value and value not in binders:
                binders.append(value)
        min_application = [layer.material.min_application_temperature for layer in layers if layer.material.min_application_temperature is not None]
        max_application = [layer.material.max_application_temperature for layer in layers if layer.material.max_application_temperature is not None]
        min_recoat = [layer.material.min_recoat_time_h for layer in layers if layer.material.min_recoat_time_h is not None]
        max_recoat = [layer.material.max_recoat_time_h for layer in layers if layer.material.max_recoat_time_h is not None]
        cure = [layer.material.full_cure_time_h for layer in layers if layer.material.full_cure_time_h is not None]
        humidity = [layer.material.max_relative_humidity for layer in layers if layer.material.max_relative_humidity is not None]
        dew_margin = [layer.material.min_dew_point_margin_c for layer in layers if layer.material.min_dew_point_margin_c is not None]
        return {
            "Связующие": ", ".join(binders) if binders else "—",
            "Мин. температура нанесения, °C": max(min_application) if min_application else None,
            "Макс. температура нанесения, °C": min(max_application) if max_application else None,
            "Мин. межслойная выдержка, ч": max(min_recoat) if min_recoat else None,
            "Макс. межслойная выдержка, ч": min(max_recoat) if max_recoat else None,
            "Полное отверждение, ч": max(cure) if cure else None,
            "Макс. относительная влажность, %": min(humidity) if humidity else None,
            "Мин. запас до точки росы, °C": max(dew_margin) if dew_margin else None,
        }
