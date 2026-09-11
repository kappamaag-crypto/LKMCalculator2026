"""Review-first importer for tabular coating-system catalogues.

Reads ``Системы 1–4`` workbooks into staging records. It never writes the
material/system database and never treats spreadsheet values as TDS or
normative evidence.
"""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
import re
from .book_index import BookSource, index_books

@dataclass(frozen=True)
class MaterialCandidate:
    name: str
    source_path: str
    sheet: str
    row_number: int
    source_sha256: str

@dataclass(frozen=True)
class SystemRowCandidate:
    source_path: str
    sheet: str
    row_number: int
    values: tuple[tuple[str, str], ...]
    material_candidates: tuple[MaterialCandidate, ...]
    source_sha256: str

def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()

def _key(value: object) -> str:
    return re.sub(r"[^a-zа-я0-9]+", "", _norm(value).lower().replace("ё", "е"))

def _material_header(key: str) -> bool:
    return any(x in key for x in ("материал", "краска", "покрытие", "грунт", "эмаль"))

def _system_header(key: str) -> bool:
    return any(x in key for x in ("система", "марка", "обозначение", "system"))

def _read(source: BookSource, root: Path) -> tuple[SystemRowCandidate, ...]:
    result: list[SystemRowCandidate] = []
    suffix = source.suffix
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Legacy XLS import requires xlrd") from exc
        book = xlrd.open_workbook(root / source.relative_path, on_demand=True)
        sheets = ((sheet.name, sheet.nrows, sheet.ncols, lambda r, s=sheet: s.row_values(r)) for sheet in book.sheets())
        for sheet_name, nrows, ncols, row_reader in sheets:
            if not nrows:
                continue
            first = row_reader(0)
            headers = tuple(_norm(v) for v in first)
            keys = tuple(_key(v) for v in first)
            material_cols = tuple(i for i, k in enumerate(keys) if _material_header(k))
            system_cols = tuple(i for i, k in enumerate(keys) if _system_header(k))
            for row_no in range(1, nrows):
                cells = tuple(_norm(v) for v in row_reader(row_no))
                pairs = tuple((headers[i] or f"column_{i + 1}", cells[i]) for i in range(min(ncols, len(cells))) if cells[i])
                if not pairs or not (system_cols or material_cols):
                    continue
                mats = tuple(MaterialCandidate(cells[i], source.relative_path, sheet_name, row_no + 1, source.sha256) for i in material_cols if i < len(cells) and cells[i])
                result.append(SystemRowCandidate(source.relative_path, sheet_name, row_no + 1, pairs, mats, source.sha256))
        book.release_resources()
        return tuple(result)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Spreadsheet import requires openpyxl") from exc
    wb = load_workbook(root / source.relative_path, read_only=True, data_only=False)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                first = next(rows)
            except StopIteration:
                continue
            headers = tuple(_norm(v) for v in first)
            keys = tuple(_key(v) for v in first)
            material_cols = tuple(i for i, k in enumerate(keys) if _material_header(k))
            system_cols = tuple(i for i, k in enumerate(keys) if _system_header(k))
            for row_no, row in enumerate(rows, 2):
                cells = tuple(_norm(v) for v in row)
                pairs = tuple((headers[i] or f"column_{i + 1}", cells[i]) for i in range(min(len(headers), len(cells))) if cells[i])
                if not pairs or not (system_cols or material_cols):
                    continue
                mats = tuple(MaterialCandidate(cells[i], source.relative_path, ws.title, row_no, source.sha256) for i in material_cols if i < len(cells) and cells[i])
                result.append(SystemRowCandidate(source.relative_path, ws.title, row_no, pairs, mats, source.sha256))
    finally:
        wb.close()
    return tuple(result)

def discover_system_catalogues(books_root: Path) -> tuple[SystemRowCandidate, ...]:
    """Discover ``Системы 1–4`` rows without database writes."""
    sources = index_books(books_root)
    selected = tuple(s for s in sources if s.suffix in {".xls", ".xlsx", ".xlsm", ".xltx", ".xltm"} and Path(s.relative_path).stem.casefold().startswith("системы"))
    result: list[SystemRowCandidate] = []
    for source in selected:
        result.extend(_read(source, books_root))
    return tuple(result)

def unique_material_candidates(rows: tuple[SystemRowCandidate, ...]) -> tuple[MaterialCandidate, ...]:
    """Return unique material candidates; persistence remains a separate review step."""
    seen: set[str] = set()
    result: list[MaterialCandidate] = []
    for row in rows:
        for candidate in row.material_candidates:
            key = re.sub(r"\s+", " ", candidate.name).strip().casefold()
            if key and key not in seen:
                seen.add(key)
                result.append(candidate)
    return tuple(result)
