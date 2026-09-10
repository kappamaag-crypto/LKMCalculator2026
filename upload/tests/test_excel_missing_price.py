"""Regression for engineering Excel export when a material has no price."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.calculator import LayerInput
from app.domain.enums import BinderType, MaterialType
from app.domain.models import Material, ObjectData
from app.infrastructure.export.excel_exporter import ExcelExporter
from app.services.calculation_service import CalculationService


def test_excel_marks_missing_price_as_unknown_without_inventing_cost(tmp_path: Path):
    material = Material(
        manufacturer="Test",
        material_name="Без цены",
        material_type=MaterialType.PRIMER_ENAMEL,
        binder_type=BinderType.EPOXY,
        density=1.4,
        solids_by_volume_percent=70.0,
        price_per_kg=None,
        price_per_liter=None,
    )
    result, validation = CalculationService().calculate_system(
        ObjectData(object_name="Тест", area_m2=100.0),
        [LayerInput(material=material, target_dft=100.0)],
    )
    assert not validation.has_errors
    assert result.total_cost_per_m2 is None
    assert result.total_cost is None

    path = tmp_path / "missing-price.xlsx"
    ExcelExporter().export_calculation(result, path)
    wb = load_workbook(path, data_only=False)

    summary = wb["Итоги"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "Стоимость ЛКМ, руб/м²" in values
    assert "—" in values

    layers = wb["Слои"]
    layer_values = [cell.value for row in layers.iter_rows() for cell in row]
    assert "Без цены" in layer_values
    assert "—" in layer_values
