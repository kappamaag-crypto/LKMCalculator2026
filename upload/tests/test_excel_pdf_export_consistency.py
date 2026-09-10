"""Cross-export regression: Excel and PDF must render the same engineering result."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from app.domain.calculator import LayerInput
from app.domain.enums import BinderType, MaterialType
from app.domain.models import Material, ObjectData
from app.infrastructure.export.excel_exporter import ExcelExporter
from app.infrastructure.export.pdf_exporter import PDFExporter
from app.services.calculation_service import CalculationService


def _result():
    materials = [
        Material(
            manufacturer="Test",
            material_name=f"Материал {i}",
            material_type=MaterialType.PRIMER_ENAMEL,
            binder_type=BinderType.EPOXY,
            density=1.4,
            solids_by_volume_percent=70.0,
            price_per_kg=100.0 + i,
        )
        for i in range(1, 5)
    ]
    result, validation = CalculationService().calculate_system(
        ObjectData(object_name="Cross export", customer="Test", area_m2=100.0),
        [LayerInput(material=m, target_dft=80.0 + i * 20.0) for i, m in enumerate(materials, 1)],
    )
    assert not validation.has_errors
    return result


def test_excel_and_pdf_share_system_result_totals(tmp_path: Path):
    result = _result()
    xlsx = tmp_path / "engineering.xlsx"
    pdf = tmp_path / "engineering.pdf"

    ExcelExporter().export_calculation(result, xlsx)
    PDFExporter().export_calculation(result, pdf)

    wb = load_workbook(xlsx, data_only=False)
    summary = wb["Итоги"]
    values = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(1, summary.max_row + 1)}
    assert values["Количество слоёв"] == len(result.layers)
    assert values["Общая толщина DFT, мкм"] == result.total_dft
    assert values["Расход ЛКМ практический, кг/м²"] == result.total_practical_consumption_kg

    reader = PdfReader(str(pdf))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "4" in text
    assert f"{result.total_dft:.0f} мкм" in text
    assert f"{result.total_practical_consumption_kg:.3f} кг/м²" in text
