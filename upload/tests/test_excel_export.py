"""Тесты Excel-экспорта."""
from __future__ import annotations
from pathlib import Path
import tempfile
import pytest
from openpyxl import load_workbook
from app.domain.models import Material, ObjectData
from app.domain.enums import MaterialType, BinderType, CorrosionCategory, DurabilityLevel
from app.domain.calculator import LayerInput
from app.services.calculation_service import CalculationService
from app.infrastructure.export.excel_exporter import ExcelExporter
from app.domain.formulas import FORMULA_VERSION

@pytest.fixture
def sample_result():
    primer = Material(material_name="Грунт-Эмаль Blank Universal", material_type=MaterialType.PRIMER_ENAMEL, binder_type=BinderType.EPOXY, density=1.4, solids_percent=73.0, solids_by_volume_percent=73.0, price_per_kg=552.0, packaging_kg=20.0, manufacturer="Blank")
    finish = Material(material_name="Эмаль Blank Finish", material_type=MaterialType.FINISH, binder_type=BinderType.POLYURETHANE, density=1.3, solids_percent=58.0, solids_by_volume_percent=58.0, price_per_kg=892.0, packaging_kg=20.0, manufacturer="Blank")
    obj = ObjectData(object_name="Тест", customer="Заказчик", area_m2=100.0, calculation_number="T-001", corrosion_category=CorrosionCategory.C4, durability=DurabilityLevel.HIGH)
    result, validation = CalculationService().calculate_system(obj, [LayerInput(material=primer, target_dft=200, losses_percent=0), LayerInput(material=finish, target_dft=100, losses_percent=0)])
    assert not validation.has_errors
    return result

def test_export_creates_file(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"; ExcelExporter().export_calculation(sample_result, path); assert path.exists() and path.stat().st_size > 1000

def test_export_has_required_sheets(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"; ExcelExporter().export_calculation(sample_result, path); names = set(load_workbook(path).sheetnames); assert {"Исходные данные","Слои","Материалы","Итоги"}.issubset(names)

def test_export_contains_date_and_formula_version(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"; ExcelExporter().export_calculation(sample_result, path); wb = load_workbook(path); ws = wb["Исходные данные"]
        values = [str(cell.value or "") for row in ws.iter_rows() for cell in row]
        assert any("Дата расчёта:" in value for value in values)
        assert any(FORMULA_VERSION in value for value in values)

def test_layers_sheet_has_data(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"; ExcelExporter().export_calculation(sample_result, path); ws = load_workbook(path)["Слои"]; assert ws.max_row >= 4; assert any(cell.value and "Blank" in str(cell.value) for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=3) for cell in row)

def test_layers_sheet_has_thinner_consumption_columns(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"; ExcelExporter().export_calculation(sample_result, path); ws = load_workbook(path)["Слои"]; headers = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]; assert "Разбавитель, кг/м²" in headers; assert "Разбавитель, л/м²" in headers; assert "Разбавитель, руб/м²" in headers

def test_materials_sheet_labels_volume_solids(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"; ExcelExporter().export_calculation(sample_result, path); ws = load_workbook(path)["Материалы"]; headers = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]; assert "Сухой остаток по объёму, %" in headers; assert "Сухой остаток, %" not in headers

def test_export_unknown_price_is_not_zero(sample_result):
    sample_result.layers[0].cost_per_m2 = None
    sample_result.total_cost_per_m2 = None
    sample_result.total_cost = None
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "unknown.xlsx"; ExcelExporter().export_calculation(sample_result, path); wb = load_workbook(path)
        ws_layers = wb["Слои"]
        assert ws_layers.cell(4, 15).value == "—"
        assert ws_layers.cell(4, 17).value == "—"
        ws_summary = wb["Итоги"]
        assert any(ws_summary.cell(row, 2).value == "—" for row in range(1, ws_summary.max_row + 1))

def test_export_has_no_procurement_or_warehouse_metrics(sample_result):
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "scope.xlsx"; ExcelExporter().export_calculation(sample_result, path); wb = load_workbook(path)
        text = " ".join(str(cell.value or "") for ws in wb.worksheets for row in ws.iter_rows() for cell in row).lower()
        forbidden = ("закуп", "остат", "склад", "упаковок", "фасов")
        assert not any(term in text for term in forbidden)

def test_export_comparison(sample_result):
    from app.domain.comparison import ComparisonEngine
    primer, finish, obj = sample_result.layers[0].material, sample_result.layers[1].material, sample_result.object_data
    sys1 = [LayerInput(material=primer, target_dft=200), LayerInput(material=finish, target_dft=100)]
    sys2 = [LayerInput(material=primer, target_dft=150), LayerInput(material=finish, target_dft=80)]
    comparison = ComparisonEngine().compare(obj, [("A", sys1), ("B", sys2)])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "cmp.xlsx"; ExcelExporter().export_comparison(comparison, path); assert path.exists(); assert "Сравнение" in load_workbook(path).sheetnames
