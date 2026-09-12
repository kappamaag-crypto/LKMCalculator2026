"""Регрессия компактного инженерного Excel-экспорта."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.calculator import LayerInput
from app.domain.comparison import ComparisonEngine
from app.domain.enums import BinderType, MaterialType
from app.domain.models import Material, ObjectData
from app.infrastructure.export.engineering_excel_exporter import EngineeringExcelExporter
from app.services.calculation_service import CalculationService


def _result(layer_count: int):
    materials = [Material(manufacturer="Blank", material_name=f"Инженерный слой {i}", material_type=MaterialType.PRIMER_ENAMEL, binder_type=BinderType.EPOXY, density=1.4, solids_by_volume_percent=70.0, price_per_kg=100.0 + i, price_per_liter=140.0 + i) for i in range(1, layer_count + 1)]
    obj = ObjectData(object_name="Тест инженерного экспорта", customer="Заказчик", area_m2=100.0)
    result, validation = CalculationService().calculate_system(obj, [LayerInput(material=m, target_dft=100 + i * 10, losses_percent=5.0) for i, m in enumerate(materials, 1)])
    assert not validation.has_errors
    return result


def _two_component_result():
    two_k = Material(manufacturer="Blank", material_name="2К Инженерный материал", material_type=MaterialType.PRIMER_ENAMEL, binder_type=BinderType.EPOXY, density=1.45, solids_by_volume_percent=72.0, price_per_kg=500.0, price_per_liter=725.0, is_two_component=True)
    finish = Material(manufacturer="Blank", material_name="Инженерный финиш", material_type=MaterialType.FINISH, binder_type=BinderType.POLYURETHANE, density=1.3, solids_by_volume_percent=60.0, price_per_kg=700.0, price_per_liter=910.0)
    result, validation = CalculationService().calculate_system(ObjectData(object_name="2К инженерный экспорт", area_m2=100.0), [LayerInput(material=two_k, target_dft=120.0, losses_percent=5.0), LayerInput(material=finish, target_dft=80.0, losses_percent=5.0)])
    assert not validation.has_errors
    return result


def test_engineering_excel_has_two_tables_and_all_layers(tmp_path: Path):
    output = tmp_path / "engineering.xlsx"; result = _result(4); EngineeringExcelExporter().export_system(result, output)
    ws = load_workbook(output, data_only=False)["Инженерный расчёт"]
    assert ws["A5"].value == "Слой"; assert [ws.cell(row, 2).value for row in range(6, 10)] == [f"Blank Инженерный слой {i}" for i in range(1, 5)]; assert ws["A10"].value == "ИТОГО"
    summary_header = 13; assert ws.cell(summary_header, 1).value == "Итог системы"
    labels = [ws.cell(row, 1).value for row in range(summary_header + 1, summary_header + 10)]
    assert "Общая толщина" in labels; assert "Общий практический расход ЛКМ" in labels; assert "Общий расход разбавителя" in labels; assert "Стоимость системы" in labels; assert "Стоимость объекта" in labels


def test_engineering_excel_does_not_expose_material_prices(tmp_path: Path):
    output = tmp_path / "engineering.xlsx"; EngineeringExcelExporter().export_system(_result(3), output); ws = load_workbook(output, data_only=False)["Инженерный расчёт"]
    values = [cell.value for row in ws.iter_rows() for cell in row]; assert not any("Цена" in str(value) for value in values if value is not None); assert any("Стоимость слоя" in str(value) for value in values if value is not None)


def test_engineering_excel_supports_two_layers(tmp_path: Path):
    output = tmp_path / "engineering.xlsx"; EngineeringExcelExporter().export_system(_result(2), output); ws = load_workbook(output, data_only=False)["Инженерный расчёт"]; assert [ws.cell(row, 2).value for row in (6, 7)] == ["Blank Инженерный слой 1", "Blank Инженерный слой 2"]


def test_comparison_excel_matches_ui_and_preserves_all_layers(tmp_path: Path):
    obj = ObjectData(object_name="Объект сравнения", area_m2=100.0); engine = ComparisonEngine(); comparison = engine.compare(obj, [("Система 4 слоя", _layers_for_export(4)), ("Система 5 слоёв", _layers_for_export(5))]); output = tmp_path / "comparison.xlsx"; EngineeringExcelExporter().export_comparison(comparison, output)
    ws = load_workbook(output, data_only=False)["Сравнение систем"]; values = [cell.value for row in ws.iter_rows() for cell in row]
    assert ws["A4"].value == "Показатель"; assert ws["B4"].value == "Система 4 слоя"; assert ws["C4"].value == "Система 5 слоёв"
    for layer_no in range(1, 6): assert f"Слой {layer_no}: материал" in values
    assert "Слой 5: разбавитель, л/м²" in values; assert "Стоимость ЛКМ, руб/м²" in values; assert "Разбавитель, руб/м²" in values; assert "ЛКМ + разбавитель, руб/м²" in values; assert ws.print_area == "'Сравнение систем'!$A$1:$C$56"; assert ws.sheet_properties.pageSetUpPr.fitToPage is True


def test_engineering_excel_2k_is_one_layer_and_does_not_expose_components(tmp_path: Path):
    output = tmp_path / "engineering-2k.xlsx"; result = _two_component_result(); EngineeringExcelExporter().export_system(result, output); ws = load_workbook(output, data_only=False)["Инженерный расчёт"]
    assert [ws.cell(row, 2).value for row in (6, 7)] == ["Blank 2К Инженерный материал", "Blank Инженерный финиш"]
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert sum("2К Инженерный материал" in str(value or "") for value in values) == 1
    assert not any(token in str(value) for value in values if value is not None for token in ("Основа", "Отвердитель", "purchase_a", "purchase_b", "sets"))


def _layers_for_export(count: int) -> list[LayerInput]:
    materials = [Material(manufacturer="Blank", material_name=f"Слой {i}", material_type=MaterialType.PRIMER_ENAMEL, binder_type=BinderType.EPOXY, density=1.4, solids_by_volume_percent=70.0, price_per_kg=100.0 + i) for i in range(1, count + 1)]
    return [LayerInput(material=m, target_dft=100 + i * 10, losses_percent=5.0) for i, m in enumerate(materials, 1)]
