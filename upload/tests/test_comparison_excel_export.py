"""Регрессия инженерного Excel-экспорта сравнения систем."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.calculator import LayerInput
from app.domain.comparison import ComparisonEngine
from app.domain.enums import BinderType, MaterialType
from app.domain.models import Material, ObjectData
from app.infrastructure.export.excel_exporter import ExcelExporter
from app.services.calculation_service import CalculationService


def _material(name: str, price: float) -> Material:
    return Material(
        manufacturer="Test",
        material_name=name,
        material_type=MaterialType.PRIMER_ENAMEL,
        binder_type=BinderType.EPOXY,
        density=1.4,
        solids_by_volume_percent=70.0,
        price_per_kg=price,
        price_per_liter=price * 1.4,
    )


def _result(layer_count: int, prefix: str):
    materials = [_material(f"{prefix} — материал {i}", 100 + i * 10) for i in range(1, layer_count + 1)]
    layers = [LayerInput(material=m, target_dft=80 + i * 20, losses_percent=5.0) for i, m in enumerate(materials)]
    result, validation = CalculationService().calculate_system(ObjectData(object_name="Тест", area_m2=100.0), layers)
    assert not validation.has_errors
    result.system.system_name = prefix
    return result


def test_comparison_excel_contains_all_layers_and_engineering_totals(tmp_path: Path):
    comparison = ComparisonEngine().compare_results(
        ObjectData(object_name="Тест", area_m2=100.0),
        [_result(2, "Система A"), _result(5, "Система B")],
    )
    path = tmp_path / "comparison.xlsx"
    ExcelExporter().export_comparison(comparison, path)

    ws = load_workbook(path, data_only=False)["Сравнение"]
    values = [str(cell.value or "") for row in ws.iter_rows() for cell in row]

    assert "Количество слоёв" in values
    assert "Общая толщина DFT, мкм" in values
    assert "Теоретический расход ЛКМ, кг/м²" in values
    assert "Практический расход ЛКМ, кг/м²" in values
    assert "Итого стоимость, руб/м²" in values
    assert "Слой 5: Материал" in values
    assert "Слой 5: DFT, мкм" in values
    assert "Слой 5: WFT, мкм" in values
    assert "Слой 5: Расход практ., кг/м²" in values
    assert "Система A" in values
    assert "Система B" in values


def test_comparison_excel_does_not_truncate_shorter_system_layers(tmp_path: Path):
    comparison = ComparisonEngine().compare_results(
        ObjectData(object_name="Тест", area_m2=100.0),
        [_result(2, "A"), _result(4, "B")],
    )
    path = tmp_path / "comparison.xlsx"
    ExcelExporter().export_comparison(comparison, path)
    ws = load_workbook(path, data_only=False)["Сравнение"]

    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Слой 4: Материал")
    assert ws.cell(row, 2).value == "—"
    assert ws.cell(row, 3).value == "Test B — материал 4"
