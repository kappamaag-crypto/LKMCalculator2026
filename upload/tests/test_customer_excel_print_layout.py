"""Регрессия печатного макета customer-facing Excel."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.calculator import LayerInput
from app.domain.enums import BinderType, MaterialType
from app.domain.models import Material, ObjectData
from app.infrastructure.export.customer_excel_exporter import CustomerExcelExporter
from app.services.calculation_service import CalculationService


def _result(layer_count: int):
    materials = [
        Material(
            manufacturer="Blank",
            material_name=f"Слой {i}",
            material_type=MaterialType.PRIMER_ENAMEL,
            binder_type=BinderType.EPOXY,
            density=1.4,
            solids_by_volume_percent=70.0,
            price_per_kg=100.0 + i,
            price_per_liter=140.0 + i,
        )
        for i in range(1, layer_count + 1)
    ]
    inputs = [
        LayerInput(material=m, target_dft=100.0 + i * 10.0, losses_percent=5.0)
        for i, m in enumerate(materials, 1)
    ]
    result, validation = CalculationService().calculate_system(
        ObjectData(object_name="Печатный тест", customer="Заказчик", area_m2=100.0),
        inputs,
    )
    assert not validation.has_errors
    assert result is not None
    return result


def _export(tmp_path: Path, layer_count: int):
    output = tmp_path / f"customer-{layer_count}.xlsx"
    exporter = CustomerExcelExporter()
    exporter.settings.report_mode = "commercial"
    exporter.export_calculation(_result(layer_count), output)
    return load_workbook(output, data_only=False)["База"]


def test_customer_excel_print_layout_two_layers(tmp_path: Path):
    ws = _export(tmp_path, 2)
    assert ws.print_area == "'База'!$B$1:$R$11"
    assert ws.print_title_rows == "1:6"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws["C11"].value == "Толщина покрытия (мкм)"


def test_customer_excel_print_layout_expands_with_layers(tmp_path: Path):
    ws = _export(tmp_path, 5)
    assert ws.print_area == "'База'!$B$1:$R$17"
    assert ws.print_title_rows == "1:6"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws["C17"].value == "Толщина покрытия (мкм)"
    assert ws["I17"].value == 650.0
